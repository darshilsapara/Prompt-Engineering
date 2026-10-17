"""Generate practical report Word document for Prompt Engineering assignment."""

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


def set_cell_shading(cell, fill_hex: str) -> None:
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill_hex)
    cell._tc.get_or_add_tcPr().append(shading)


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    doc.add_heading(text, level=level)


def add_screenshot(doc: Document, image_path: Path, caption: str, width_in: float = 5.5) -> None:
    """Insert a centered screenshot with caption if the file exists."""
    doc.add_paragraph(caption)
    if image_path.is_file():
        pic = doc.add_picture(str(image_path), width=Inches(width_in))
        pic_paragraph = doc.paragraphs[-1]
        pic_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    else:
        doc.add_paragraph(f"[Missing image: {image_path.name}]")
    doc.add_paragraph()


def add_table_with_header(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = h
        set_cell_shading(hdr[i], "D9E2F3")
        for p in hdr[i].paragraphs:
            for run in p.runs:
                run.bold = True
    for r_idx, row in enumerate(rows):
        for c_idx, val in enumerate(row):
            table.rows[r_idx + 1].cells[c_idx].text = val
    doc.add_paragraph()


def build_document() -> Document:
    base_dir = Path(__file__).resolve().parent
    shots = base_dir / "docs" / "screenshots"

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    # Cover / Title
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(
        "Development of NeonRacer Motor Racing Game\n"
        "Using Generative AI and Prompt Engineering"
    )
    run.bold = True
    run.font.size = Pt(18)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.add_run(f"\nPractical Report\nSubmitted: {date.today().strftime('%d %B %Y')}\n")
    doc.add_page_break()

    # 1. Aim
    add_heading(doc, "1. Aim", 1)
    doc.add_paragraph(
        "The aim of this practical is to design, develop, and document a simple 2D browser-based "
        "motor racing game (NeonRacer) using prompt engineering with a generative AI assistant. "
        "The work demonstrates how structured prompts can be used for game planning, modular "
        "architecture design, code generation, testing, debugging, and feature enhancement—while "
        "producing a working application with Django, HTML5 Canvas, Vue.js, and SQLAlchemy."
    )

    # 2. Software / Tools
    add_heading(doc, "2. Software / Tools Used", 1)
    tools = [
        "Python 3.x — server-side logic and API",
        "Django 6.x — web framework, routing, templates",
        "HTML5 Canvas + JavaScript — 2D game rendering and game loop",
        "Vue.js 3 — user interface (HUD, menus, leaderboard)",
        "SQLAlchemy + SQLite — high-score persistence (racer_game.db)",
        "Alembic — database schema migrations",
        "CSS3 — neon-themed UI styling",
        "Cursor IDE with AI (Composer) — prompt-driven development",
        "Git / GitHub — version control and project submission",
        "Microsoft Word — practical documentation",
    ]
    for t in tools:
        doc.add_paragraph(t, style="List Bullet")

    # 3. Game Description
    add_heading(doc, "3. Game Description", 1)

    add_heading(doc, "3.1 Objective", 2)
    doc.add_paragraph(
        "The player drives a car on a three-lane neon highway. The objective is to survive as long "
        "as possible by dodging opponent vehicles, collecting coins for bonus points, and gathering "
        "fuel to stay alive. Higher survival time and successful dodges increase the score. "
        "Top scores are saved on a server-side leaderboard."
    )

    add_heading(doc, "3.2 Rules", 2)
    rules = [
        "The player may occupy only one of three lanes at a time (left, centre, right).",
        "Opponent cars spawn at the top and move toward the player; contact causes a crash (game over).",
        "Fuel drains continuously during play; reaching zero fuel ends the game.",
        "Coins and fuel pickups appear on lanes; collecting them adds score or restores fuel.",
        "Game speed and difficulty increase as level rises (based on score).",
        "Each completed run can be submitted with a player name to the high-score API.",
    ]
    for r in rules:
        doc.add_paragraph(r, style="List Bullet")

    add_heading(doc, "3.3 Controls", 2)
    doc.add_paragraph("Keyboard: Left Arrow — move one lane left; Right Arrow — move one lane right.")
    doc.add_paragraph("Mobile / touch: On-screen LEFT and RIGHT buttons when playing.")
    doc.add_paragraph("Menus: START RACE, RACE AGAIN, LEADERBOARD, Main Menu buttons.")

    add_heading(doc, "3.4 Game Objects", 2)
    objects = [
        "Player car — user-controlled vehicle at bottom of road",
        "Opponent cars — five colour palettes, spawn in random lanes",
        "Coins — bonus score pickups",
        "Fuel canisters — restore fuel gauge",
        "Road — three lanes with animated dashed dividers",
        "Particles — explosion effect on crash or pickup",
        "HUD elements — score, level, fuel bar, coins, speed",
    ]
    for o in objects:
        doc.add_paragraph(o, style="List Bullet")

    add_heading(doc, "3.5 Score System", 2)
    doc.add_paragraph(
        "Score increases from: (1) survival time (continuous per-frame rate scaled by speed), "
        "(2) +10 points per opponent car safely passed (dodge bonus), "
        "(3) +50 points per coin collected. Level increases every 800 score points."
    )

    add_heading(doc, "3.6 Game-Over Conditions", 2)
    doc.add_paragraph(
        "Game over occurs when: (1) the player collides with an opponent car (AABB collision), or "
        "(2) fuel reaches zero. A crash animation plays, final statistics are shown, and the score "
        "may be saved to the leaderboard."
    )

    add_heading(doc, "3.7 Restart Functionality", 2)
    doc.add_paragraph(
        "From the Game Over screen, RACE AGAIN resets game state and starts a new session. "
        "From Leaderboard or Game Over, Main Menu returns to the start screen with name input "
        "and top scores preview."
    )

    # 4. Game Flow
    add_heading(doc, "4. Game Flow", 1)
    doc.add_paragraph(
        "The following flowchart describes the logical flow of NeonRacer from launch to restart."
    )
    flow = """
START
  ↓
Load Start Screen (fetch leaderboard from API)
  ↓
Player enters name → Press START RACE
  ↓
Initialize GameEngine (reset score, fuel, level, entities)
  ↓
┌──────────────────────────────────────┐
│         MAIN GAME LOOP (RAF)          │
│  Read keyboard / mobile input           │
│  Smooth lane transition for player      │
│  Update score, level, speed, fuel       │
│  Spawn / move opponents and pickups     │
│  Particle updates                       │
│  Collision: player vs opponent?         │
│  Fuel <= 0?                             │
└──────────────────────────────────────┘
  ↓                              ↓
 YES (collision / no fuel)      NO → continue loop
  ↓
Crash particles → Game Over screen
  ↓
POST score to API → refresh leaderboard
  ↓
RACE AGAIN? ──YES──→ Initialize GameEngine
  ↓ NO
Main Menu or END
"""
    p = doc.add_paragraph(flow)
    p.style = "No Spacing"
    for run in p.runs:
        run.font.name = "Consolas"
        run.font.size = Pt(9)

    doc.add_page_break()

    # STEP 1 – Game Design (Prompt 1 output)
    add_heading(doc, "5. STEP 1 – Game Design (Prompt 1 – Game Planning)", 1)

    add_heading(doc, "5.1 Game Objective", 2)
    doc.add_paragraph(
        "Survive on a three-lane highway, maximize score through time, dodges, and coins, "
        "and compete for leaderboard ranking."
    )

    add_heading(doc, "5.2 Player Controls", 2)
    doc.add_paragraph("Arrow Left / Arrow Right for lane changes; optional touch buttons.")

    add_heading(doc, "5.3 Variables Required", 2)
    add_table_with_header(
        doc,
        ["Variable / State", "Purpose"],
        [
            ["player.lane, player.x, player.y", "Player position and target lane"],
            ["score, scoreFloat", "Integer display score and fractional accumulator"],
            ["level, speed, frame", "Difficulty progression"],
            ["fuel", "Resource; drains per frame"],
            ["coins, carsDodged", "Statistics and bonuses"],
            ["opponents[]", "Active enemy cars"],
            ["pickups[]", "Coins and fuel on road"],
            ["particles[]", "Visual effects"],
            ["gameState (Vue)", "start | playing | gameover | leaderboard"],
            ["leaderboard (Vue)", "Top scores from API"],
        ],
    )

    add_heading(doc, "5.4 Functions / Components Required", 2)
    add_table_with_header(
        doc,
        ["Function / Class", "Purpose"],
        [
            ["GameEngine.start() / stop()", "Lifecycle control"],
            ["GameEngine._loop()", "Main animation frame loop"],
            ["GameEngine._update()", "Logic: input, movement, spawn, collision"],
            ["GameEngine._render()", "Draw road, entities, HUD canvas layer"],
            ["rectsOverlap()", "AABB collision detection"],
            ["drawCar(), drawCoin(), drawFuelPickup()", "Rendering helpers"],
            ["fetchLeaderboard(), submitScore()", "Vue API integration"],
            ["get_scores(), post_score() (Django)", "REST backend for scores"],
        ],
    )

    add_heading(doc, "5.5 Collision Conditions", 2)
    doc.add_paragraph(
        "Player–opponent: axis-aligned bounding box overlap with small margin "
        "(rectsOverlap in game.js). Player–pickup: circle distance check between "
        "player centre and pickup centre. No collision between opponents and pickups required."
    )

    add_heading(doc, "5.6 Score Calculation", 2)
    doc.add_paragraph(
        "TIME_SCORE_RATE × speed per frame (floored into score); +DODGE_BONUS (10) per car passed; "
        "+COIN_VALUE (50) per coin. Level = 1 + floor(score / 800)."
    )

    add_heading(doc, "5.7 Suggested Program Structure", 2)
    doc.add_paragraph(
        "Frontend: static/js/game.js (GameEngine, Road, Particle), "
        "game/templates/game/index.html (Vue app + canvas), static/css/style.css. "
        "Backend: game/views.py, game/models.py (SQLAlchemy HighScore), game/db.py, game/urls.py."
    )

    # STEP 2 – Architecture
    add_heading(doc, "6. STEP 2 – Program Architecture (Prompt 2)", 1)
    arch_rows = [
        ("Game initialization", "GameEngine constructor, _resetState(), Django init_db()", "player, speed, fuel, lists"),
        ("Player creation", "_resetState sets lane 1, position bottom", "player object"),
        ("Keyboard input", "_bindKeys(), _movePlayer()", "keyLeft, keyRight"),
        ("Object generation", "_spawnOpponent(), _spawnPickup()", "oppTimer, pickupTimer"),
        ("Movement", "_update() updates y positions and lane easing", "speed, opponents, pickups"),
        ("Collision detection", "rectsOverlap(), distance for pickups", "margin, radii"),
        ("Score management", "_update() + callbacks onScore/onCoins", "score, carsDodged"),
        ("Game-over detection", "collision or fuel <= 0 → _triggerCrash()", "crashed, onGameOver"),
        ("Game restart", "Vue startGame() destroys engine, new GameEngine", "gameState"),
        ("Main game loop", "_loop() via requestAnimationFrame", "running flag"),
    ]
    add_table_with_header(
        doc,
        ["Component", "Purpose", "Key variables"],
        [[a, b, c] for a, b, c in arch_rows],
    )

    # STEP 3 – Code generation summary
    add_heading(doc, "7. STEP 3 – Code Generation Summary (Prompt 3)", 1)

    add_heading(doc, "7.1 How to Run the Project", 2)
    steps = [
        "Install dependencies: pip install -r requirements.txt",
        "Apply DB migration (optional fresh DB): alembic upgrade head",
        "Start server: python manage.py runserver",
        "Open browser: http://127.0.0.1:8000/",
        "Run automated tests: python manage.py test game",
    ]
    for i, s in enumerate(steps, 1):
        doc.add_paragraph(f"{i}. {s}")

    add_heading(doc, "7.2 Required Libraries", 2)
    doc.add_paragraph("Django, SQLAlchemy, Alembic (see requirements.txt). Browser loads Vue 3 from CDN.")

    add_heading(doc, "7.3 Important Functions", 2)
    doc.add_paragraph(
        "GameEngine._update() — core rules; rectsOverlap() — collisions; "
        "views.post_score / get_scores — persistence; Vue startGame / onGameOver — UI flow."
    )

    add_heading(doc, "7.4 How Collision Detection Works", 2)
    doc.add_paragraph(
        "Opponents use AABB: compare centre positions and half-widths/heights with a margin. "
        "Pickups use Euclidean distance between centres compared to combined radius."
    )

    add_heading(doc, "7.5 GitHub Repository", 2)
    doc.add_paragraph(
        "[Student to insert GitHub repository URL after upload. Example: "
        "https://github.com/<username>/neonracer-prompt-engineering]"
    )

    doc.add_page_break()

    # STEP 4 – Testing
    add_heading(doc, "8. STEP 4 – Test Cases", 1)
    doc.add_paragraph(
        "Manual and automated tests performed on the implemented game. "
        "Status reflects verification during development."
    )
    add_table_with_header(
        doc,
        ["Test No.", "Test Performed", "Expected Result", "Actual Result", "Status"],
        [
            ["1", "Start game (START RACE)", "Game starts, HUD visible, canvas animates", "Road scrolls, player car shown", "Pass"],
            ["2", "Press Left Arrow", "Player moves to left lane", "Lane changes with smooth easing", "Pass"],
            ["3", "Press Right Arrow", "Player moves to right lane", "Lane changes correctly", "Pass"],
            ["4", "Collision with opponent", "Game Over, explosion particles", "Crash flash, stats screen", "Pass"],
            ["5", "Survive / collect coin", "Score increases", "Score updates in HUD", "Pass"],
            ["6", "RACE AGAIN after game over", "New game starts", "Score reset, play resumes", "Pass"],
            ["7", "Fuel depletion", "Game Over when fuel empty", "Game ends without collision", "Pass"],
            ["8", "GET /api/scores/", "JSON list of scores", "Returns scores array", "Pass"],
            ["9", "POST /api/scores/submit/", "Score saved", "201 success, entry in DB", "Pass"],
            ["10", "python manage.py test game", "All tests pass", "4 tests OK", "Pass"],
        ],
    )

    # STEP 5 – Debugging
    add_heading(doc, "9. STEP 5 – Debugging (Prompt 4 Example)", 1)
    doc.add_paragraph("Problem: Home page failed when rendered by Django; tests reported TemplateSyntaxError.")
    doc.add_paragraph("Error: Could not parse the remainder: '()' from 'score.toLocaleString()'")
    doc.add_paragraph("Expected: Vue bindings should display score on the page.")
    doc.add_paragraph("Actual: Django template engine interpreted {{ }} as Django tags.")
    doc.add_paragraph(
        "Cause: Django and Vue both use double curly braces. "
        "Fix: Wrapped the Vue root markup in {% verbatim %}…{% endverbatim %} in index.html "
        "so Django leaves Vue syntax unchanged. Unrelated game logic was not modified."
    )

    # STEP 6 – Improvements
    add_heading(doc, "10. STEP 6 – Feature Enhancements (Prompt 5)", 1)
    doc.add_paragraph("At least three improvements were implemented beyond a minimal dodging game:")
    improvements = [
        ("High-score system", "Django API + SQLAlchemy + leaderboard UI on start/game over screens"),
        ("Start screen & Game-over screen", "Vue overlays with name input, stats grid, navigation"),
        ("Increasing difficulty", "Level-based speed and spawn interval; speed lines at high speed"),
        ("Better graphics", "Neon cars, gradients, particle explosions, glassmorphism HUD"),
        ("Multiple obstacle types", "Five opponent colour palettes"),
        ("Fuel & coin systems", "Secondary game-over condition and bonus scoring"),
        ("Mobile controls", "Touch LEFT/RIGHT buttons"),
        ("Alembic migrations", "Schema versioning for high_scores table"),
    ]
    add_table_with_header(
        doc,
        ["Feature", "Description"],
        improvements,
    )

    doc.add_page_break()

    # PART E – Prompting techniques
    add_heading(doc, "11. PART E – Prompt Engineering Techniques Used", 1)
    doc.add_paragraph("At least five techniques were applied during development:")
    add_table_with_header(
        doc,
        ["No.", "Technique", "How It Was Used in This Project"],
        [
            ["1", "Role prompting", "Asked AI to act as game developer to plan NeonRacer mechanics and stack"],
            ["2", "Context provision", "Provided Django + Vue + existing file structure when completing project"],
            ["3", "Task decomposition", "Separated planning, architecture, backend API, canvas engine, and UI"],
            ["4", "Constraint prompting", "Specified arrow keys only, no removal of score system when adding features"],
            ["5", "Structured output", "Requested game design sections (rules, variables, collisions) before code"],
            ["6", "Iterative prompting", "Basic loop first, then leaderboard, fuel, particles, styling"],
            ["7", "Debugging prompting", "Supplied TemplateSyntaxError, expected vs actual for Vue/Django fix"],
            ["8", "Refinement", "Asked to complete only missing parts (Alembic, tests, template fix)"],
            ["9", "Verification", "Ran manage.py test and API checks against requirements"],
        ],
    )

    # PART G – Prompts Used
    add_heading(doc, "12. PART G – Prompts Used (Sample)", 1)
    prompts = [
        (
            "Game planning",
            "I want to develop a simple 2D NeonRacer motor racing game using Django, HTML5 Canvas, "
            "and JavaScript. Before code, provide game objective, rules, controls, objects, "
            "variables, functions, collisions, score, game-over, and program structure. Do not code yet.",
            "Full game design document with lane-based dodging, fuel, and scoring.",
            "Adjusted level gap to 800 points and three lanes for playability.",
        ),
        (
            "Program architecture",
            "Based on the design, create modular structure: init, player, input, spawn, movement, "
            "collision, score, game-over, restart, main loop. No full code yet.",
            "Component table mapping to GameEngine methods and Vue callbacks.",
            "Merged API layer as separate Django views component.",
        ),
        (
            "Code generation",
            "Generate complete working code: meaningful names, functions, comments, keyboard controls, "
            "collision, score, game over, restart. Explain how to run and how score works.",
            "game.js GameEngine, index.html Vue app, views and SQLAlchemy models.",
            "Tested on runserver; fixed static files path in settings.",
        ),
        (
            "Complete incomplete work",
            "Complete this project; do only things which are not complete.",
            "Alembic setup, tests, EMAIL settings, verbatim Vue fix.",
            "Accepted changes; prepared GitHub upload separately.",
        ),
        (
            "Documentation",
            "Make Word document for all practical requirements including prompt log and screenshots.",
            "This report document.",
            "Embedded screenshots in Part I of this report.",
        ),
        (
            "Debugging",
            "TemplateSyntaxError on score.toLocaleString() — expected Vue rendering, Django parsed tags.",
            "Identified brace conflict; verbatim block solution.",
            "Applied {% verbatim %} in index.html.",
        ),
    ]
    for i, (purpose, prompt, result, mod) in enumerate(prompts, 1):
        add_heading(doc, f"12.{i} Prompt — {purpose}", 3)
        doc.add_paragraph("Purpose: " + purpose)
        doc.add_paragraph("Prompt used:")
        doc.add_paragraph(prompt)
        doc.add_paragraph("AI response / result: " + result)
        doc.add_paragraph("Modification by student: " + mod)

    doc.add_page_break()

    # PART H – Prompt Log
    add_heading(doc, "13. PART H – Prompt Log", 1)
    add_table_with_header(
        doc,
        ["Prompt No.", "Purpose", "Prompting Technique", "AI Output", "Modification Made"],
        [
            ["1", "Game planning", "Task decomposition", "Game design for NeonRacer", "Tuned rules and scoring constants"],
            ["2", "Program structure", "Structured output", "Module/function list", "Aligned with GameEngine class"],
            ["3", "Code generation", "Constraint prompting", "Full stack initial code", "Tested and integrated"],
            ["4", "Debugging", "Context + error prompting", "Vue/Django template fix", "Applied verbatim block"],
            ["5", "Feature addition", "Iterative prompting", "Leaderboard, fuel, UI screens", "Kept arrow-key controls"],
            ["6", "Code review", "Verification", "Tests + Alembic + settings", "Ran manage.py test game"],
        ],
    )

    # PART I – Screenshots
    add_heading(doc, "14. PART I – Screenshots", 1)
    doc.add_paragraph(
        "Screenshots captured from the running NeonRacer application "
        "(python manage.py runserver → http://127.0.0.1:8000/)."
    )
    add_screenshot(
        doc,
        shots / "01_start_screen.png",
        "Figure 1 — Start screen: title, racer name, START RACE, controls hint, top racers preview.",
    )
    add_screenshot(
        doc,
        shots / "02_gameplay.png",
        "Figure 2 — Running game (gameplay): three-lane road, player car, opponents, fuel pickup.",
    )
    add_screenshot(
        doc,
        shots / "03_game_over.png",
        "Figure 3 — Score / achievement state and Game Over (after collision): final score, "
        "statistics, new high score flag, and restart options.",
    )
    add_screenshot(
        doc,
        shots / "04_leaderboard.png",
        "Figure 4 — Leaderboard: top racers with scores, levels, and coins from the API.",
    )

    add_heading(doc, "15. Conclusion", 1)
    doc.add_paragraph(
        "NeonRacer demonstrates end-to-end development with generative AI and deliberate prompt "
        "engineering: from structured game design through modular implementation, testing, debugging, "
        "and iterative enhancements. The final product meets assignment requirements including "
        "player control, obstacles, collision, scoring, game over, restart, UI screens, high scores, "
        "and documented prompt techniques."
    )

    add_heading(doc, "16. Declaration", 1)
    doc.add_paragraph(
        "I confirm that this report documents my practical work, including prompts used with AI "
        "assistance, testing performed, and modifications made by me. GitHub source code link: "
        "[to be filled]."
    )
    doc.add_paragraph("\nStudent Name: _________________________    Roll No.: _________________________")
    doc.add_paragraph("Signature: _________________________    Date: _________________________")

    return doc


def main() -> None:
    out = Path(__file__).resolve().parent / "NeonRacer_Prompt_Engineering_Practical_Report.docx"
    doc = build_document()
    doc.save(out)
    print(f"Saved: {out}")


if __name__ == "__main__":
    main()
